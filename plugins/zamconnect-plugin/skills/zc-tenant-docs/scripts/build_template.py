"""
One-time (re-runnable) builder for the api-spec-sync skill's reusable Word template.

Derives templates/api-specification-template.docx from the existing ZIMS specification, so
styles, numbering, the cover-page logo, headers/footers, and the Table of Contents field all
carry over - but replaces every tenant-specific literal (title, role, version, date) with a
{{TOKEN}}, strips all body content after the Document History table (the renderer fills that
in per-run), and fixes the heading-list indent bug at the template level instead of patching it
on every build.

Run this again only if the template itself needs to change (e.g. a new placeholder, a styling
fix). Per-tenant regeneration never touches this script - see scripts/openapi_to_docx.py.
"""
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # .claude/skills/api-spec-sync
REPO_ROOT = os.path.abspath(os.path.join(SKILL_DIR, "..", "..", ".."))

SOURCE = os.path.join(REPO_ROOT, "ZIMS (c) API Specification.docx")
OUTPUT = os.path.join(SKILL_DIR, "templates", "api-specification-template.docx")


def set_para_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


shutil.copyfile(SOURCE, OUTPUT)
d = docx.Document(OUTPUT)
body = d.element.body
children = list(body.iterchildren())

# --- cover page --------------------------------------------------------------------------
# Index map verified against this exact source file (see build_apis_docx.py, same technique
# used successfully twice this session).
COVER = {
    0: "{{TITLE}}",
    1: "({{TENANT}})",
    2: "({{ROLE}})",
    3: "API Specification",
    5: "Version {{VERSION}} ● {{DATE}}",
    22: "Last edited: {{DATE}}",
}
for idx, text in COVER.items():
    set_para_text(Paragraph(children[idx], d), text)

# --- document history table ---------------------------------------------------------------
hist = docx.table.Table(children[34], d)
while len(hist.rows) > 2:
    hist._tbl.remove(hist.rows[-1]._tr)
placeholders = ["{{HISTORY_DESCRIPTION}}", "{{HISTORY_AUTHOR}}", "{{HISTORY_VERSION}}", "{{HISTORY_DATE}}"]
for cell, value in zip(hist.rows[1].cells, placeholders):
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para_text(cell.paragraphs[0], value)

# --- drop everything after the history table (renderer fills this in per run) -------------
for child in children[35:]:
    if child.tag != qn("w:sectPr"):
        body.remove(child)

# --- header / footer: precise, run-verified token substitution ----------------------------
# Run boundaries confirmed by direct inspection (python-docx splits "Version 1.9" and the
# copyright year into their own runs due to prior track-changes/find-replace history in the
# source document) - see the empirical dump this was built from, not guessed.
for section in d.sections:
    for container in (section.header, section.footer,
                      section.first_page_header, section.first_page_footer,
                      section.even_page_header, section.even_page_footer):
        if container is None:
            continue
        for p in container.paragraphs:
            for r in p.runs:
                if "ZIMS" in r.text:
                    r.text = r.text.replace("ZIMS", "{{TENANT}}")
                if "(Consumer)" in r.text:
                    r.text = r.text.replace("(Consumer)", "({{ROLE}})")
                # Header: "Version 1." + "9" as two adjacent runs.
                if re.fullmatch(r"Version 1\.", r.text) and "Version" in p.text:
                    r.text = "Version {{VERSION_MAJOR}}."
                if re.fullmatch(r"\s*9\s*", r.text) and "Version" in p.text:
                    r.text = r.text.replace("9", "{{VERSION_MINOR}}")
                if "May 5, 2026" in r.text:
                    r.text = r.text.replace("May 5, 2026", "{{DATE}}")
                # Footer copyright year: only the bare "2026" run, never the one inside the
                # "Last modified: May 5, 2026" date (already handled above).
                if re.fullmatch(r"2026", r.text.strip()):
                    r.text = r.text.replace("2026", "{{YEAR}}")

d.save(OUTPUT)

# --- bake style-level fixes into the template permanently ----------------------------------
# These are template properties, not per-run content, so they belong here once rather than
# being re-patched by every tenant's render (which is what this session's one-off APIS script
# did, out of necessity, before this template existed).
with zipfile.ZipFile(OUTPUT) as zf:
    parts = {n: zf.read(n) for n in zf.namelist()}


def el(tag, **attrs):
    from docx.oxml import OxmlElement
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


# 1) Heading-list indent bug: abstractNumId 18 (via numId 1) ships with a ~4 inch left indent
#    on level 2 (ilvl=1), pushing every Heading 3 into the middle of the page.
numbering = parts["word/numbering.xml"].decode("utf-8")
numbering, subs = re.subn(r'<w:ind w:left="5742" w:hanging="432"/>',
                          '<w:ind w:left="720" w:hanging="720"/>', numbering)
assert subs == 1, f"expected exactly one level-2 indent to patch, found {subs}"
parts["word/numbering.xml"] = numbering.encode("utf-8")

# 2) Body font: Times New Roman 12 (theme + docDefaults + Normal/No Spacing/List Paragraph).
#    Heading sizes are left alone (already Times New Roman 22pt/16pt in the source) to preserve
#    the visual hierarchy against the now-12pt body.
theme = parts["word/theme/theme1.xml"].decode("utf-8")
theme = re.sub(r'(<a:(?:majorFont|minorFont)>\s*<a:latin typeface=")[^"]*"',
               r'\1Times New Roman"', theme)
parts["word/theme/theme1.xml"] = theme.encode("utf-8")

styles_xml = parts["word/styles.xml"].decode("utf-8")
styles_xml = styles_xml.replace(
    '<w:rFonts w:asciiTheme="minorHAnsi" w:eastAsiaTheme="minorHAnsi" '
    'w:hAnsiTheme="minorHAnsi" w:cstheme="minorBidi"/><w:sz w:val="22"/><w:szCs w:val="22"/>',
    '<w:rFonts w:ascii="Times New Roman" w:eastAsia="Times New Roman" '
    'w:hAnsi="Times New Roman" w:cs="Times New Roman"/><w:sz w:val="24"/><w:szCs w:val="24"/>')

NEW_RPR_FONTS = ('<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" '
                 'w:eastAsia="Times New Roman" w:cs="Times New Roman"/>'
                 '<w:sz w:val="24"/><w:szCs w:val="24"/>')

for style_id in ("Normal", "NoSpacing", "ListParagraph"):
    m = re.search(rf'<w:style [^>]*w:styleId="{style_id}".*?</w:style>', styles_xml, re.S)
    if not m:
        continue
    block = m.group(0)

    # Normal/NoSpacing in the ZIMS source already carry their own <w:rFonts>/<w:sz> overrides
    # (e.g. an explicit Calibri). Simply prepending our Times New Roman rFonts left the original
    # as a second, later sibling inside the same <w:rPr> - and a later duplicate element wins in
    # OOXML, so the doc silently rendered in Calibri despite this patch "succeeding". Strip any
    # pre-existing rFonts/sz/szCs before inserting ours so there is no duplicate left to win.
    rpr_match = re.search(r"<w:rPr>(.*?)</w:rPr>|<w:rPr/>", block, re.S)
    if rpr_match is None:
        new_block = block[:-len("</w:style>")] + f"<w:rPr>{NEW_RPR_FONTS}</w:rPr></w:style>"
    else:
        inner = rpr_match.group(1) or ""
        inner = re.sub(r"<w:rFonts\b[^>]*/>", "", inner)
        inner = re.sub(r"<w:sz\b[^>]*/>", "", inner)
        inner = re.sub(r"<w:szCs\b[^>]*/>", "", inner)
        new_rpr = f"<w:rPr>{NEW_RPR_FONTS}{inner}</w:rPr>"
        new_block = block[:rpr_match.start()] + new_rpr + block[rpr_match.end():]

    styles_xml = styles_xml[:m.start()] + new_block + styles_xml[m.end():]

# 3) Page break before every top-level (Heading 2) section.
h2 = re.search(r'<w:style [^>]*w:styleId="Heading2".*?</w:style>', styles_xml, re.S).group(0)
if "pageBreakBefore" not in h2:
    h2_fixed = re.sub(r'(<w:pPr>)', r'\1<w:pageBreakBefore/>', h2, count=1)
    styles_xml = styles_xml.replace(h2, h2_fixed, 1)

parts["word/styles.xml"] = styles_xml.encode("utf-8")

with zipfile.ZipFile(OUTPUT, "w", zipfile.ZIP_DEFLATED) as zf:
    for name, data in parts.items():
        zf.writestr(name, data)

print(f"Template written: {OUTPUT}")

# --- verify ---------------------------------------------------------------------------------
check = docx.Document(OUTPUT)
print(f"paragraphs: {len(check.paragraphs)}  tables: {len(check.tables)}")
for p in check.paragraphs[:10]:
    if p.text.strip():
        print("  cover:", repr(p.text))
for row in check.tables[0].rows:
    print("  history row:", [c.text for c in row.cells])
print("header:", [p.text for p in check.sections[0].header.paragraphs if p.text.strip()])
print("footer:", [p.text for p in check.sections[0].footer.paragraphs if p.text.strip()])

with zipfile.ZipFile(OUTPUT) as zf:
    ok = zf.testzip() is None
    for n in zf.namelist():
        if n.endswith((".xml", ".rels")):
            from lxml import etree
            try:
                etree.fromstring(zf.read(n))
            except Exception as e:
                ok = False
                print("XML ERROR", n, e)
    st = zf.read("word/styles.xml").decode("utf-8")
    print("package valid:", ok)
    print("body font TNR 12:", 'w:ascii="Times New Roman"' in st and '<w:sz w:val="24"/>' in st)
    print("page break before H2:", "pageBreakBefore" in re.search(
        r'w:styleId="Heading2".*?</w:style>', st, re.S).group(0))
